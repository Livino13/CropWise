"""Generate the CropWise Version 2.0 change report.

This is intentionally separate from generate_report.py / the main SRS. It records
the changes implemented so far, the updated software and hardware requirements,
the current architecture, the methodology, and the verification performed.
"""
import os
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

OUT = os.path.join(
    os.path.dirname(os.path.abspath(__file__)),
    "CropWise_Version_2_Change_Report.docx",
)

doc = Document()
doc.styles["Normal"].font.name = "Calibri"
doc.styles["Normal"].font.size = Pt(10.5)
for section in doc.sections:
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

doc.core_properties.title = "CropWise Version 2.0 Change Report"
doc.core_properties.author = "Liv Major Project"
doc.core_properties.subject = "Changes, requirements, architecture, methodology and verification"
doc.core_properties.comments = "Separate change report; the main SRS is unchanged."


def centered(text, size=12, bold=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.bold = bold
    r.font.size = Pt(size)
    return p


def para(text, bold_lead=None):
    p = doc.add_paragraph()
    if bold_lead:
        r = p.add_run(bold_lead)
        r.bold = True
    p.add_run(text)
    return p


def bullet(bold_part, rest=""):
    p = doc.add_paragraph(style="List Bullet")
    if bold_part:
        r = p.add_run(bold_part)
        r.bold = True
    if rest:
        p.add_run(rest)
    return p


def code(text):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(0.3)
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(4)
    r = p.add_run(text)
    r.font.name = "Consolas"
    r.font.size = Pt(8.5)
    r.font.color.rgb = RGBColor(0x22, 0x22, 0x22)
    return p


def table(headers, rows, font_size=8.5):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Light Grid Accent 1"
    t.autofit = True
    for i, h in enumerate(headers):
        t.rows[0].cells[i].text = h
    for row_data in rows:
        row = t.add_row()
        for i, val in enumerate(row_data):
            row.cells[i].text = str(val)
    for row in t.rows:
        for cell in row.cells:
            for p in cell.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(font_size)
    for p in t.rows[0].cells[0].paragraphs[0].runs:
        p.bold = True
    return t


# ============================================================
# TITLE PAGE
# ============================================================
for _ in range(4):
    doc.add_paragraph("")
centered("CROPWISE", 20, True)
centered("VERSION 2.0 CHANGE REPORT", 15, True)
doc.add_paragraph("")
centered("Changes, Updated Requirements, Architecture, Methodology and Verification", 12)
doc.add_paragraph("")
centered("Liv Major Project", 11)
centered("22 October 2026", 11)
doc.add_paragraph("")
centered("Backend: FastAPI + Python   |   Frontend: React 18 + Vite", 9)
centered("Intelligence: FCM + PSO + Random Forest  |  Deterministic agronomy engine  |  Gemini assistant", 9)
doc.add_paragraph("")
centered("Separate change document - the main SRS remains unchanged", 9, True)
doc.add_page_break()

# ============================================================
# DOCUMENT CONTROL
# ============================================================
doc.add_heading("Document Control", level=1)
table(
    ["Field", "Value"],
    [
        ("Document", "CropWise Version 2.0 Change Report"),
        ("Version", "2.0"),
        ("Date", "22 October 2026"),
        ("Project", "Liv Major"),
        ("Status", "Reflects the current working tree; changes are not committed or pushed"),
        ("Scope", "Changes made since the original crop recommendation report, plus current requirements, architecture and methodology"),
        ("Main SRS", "CropWise_System_Requirements_Specification.docx (not modified by this document)"),
        ("Source of truth", "Backend and frontend source code, model artifact, and verification run on 22 October 2026"),
    ],
)
doc.add_paragraph("")
para(
    "This document is intentionally separate from the main System Requirements and Specification. "
    "It records what changed in Version 2.0, why the system is now a hybrid advisory platform, "
    "and what the current implementation requires and verifies. It does not replace the original SRS."
)
doc.add_paragraph("How to read this report:")
bullet("Sections 1-2: ", "executive summary and the complete change log.")
bullet("Sections 3-4: ", "current software and hardware requirements.")
bullet("Sections 5-6: ", "current architecture and methodology.")
bullet("Sections 7-9: ", "verification results, limitations, and deployment notes.")
bullet("Appendices: ", "API summary, model metadata, environment variables, file map and glossary.")
doc.add_page_break()

# ============================================================
# CONTENTS
# ============================================================
doc.add_heading("Contents", level=1)
toc = [
    ("1.", "Executive Summary", True),
    ("2.", "Changes Made in Version 2.0", True),
    ("2.1", "Frontend and User Experience", False),
    ("2.2", "Backend, API and Services", False),
    ("2.3", "Machine Learning Model", False),
    ("2.4", "Database and Data Handling", False),
    ("2.5", "Navigation, Content and Compatibility", False),
    ("3.", "Updated Software Requirements", True),
    ("4.", "Updated Hardware Requirements", True),
    ("5.", "Current System Architecture", True),
    ("5.1", "Architectural Overview", False),
    ("5.2", "Layer Responsibilities", False),
    ("5.3", "Data Flows", False),
    ("5.4", "Caching, Failure Handling and Security", False),
    ("6.", "Current Methodology", True),
    ("6.1", "Development Methodology", False),
    ("6.2", "Machine Learning Methodology", False),
    ("6.3", "Agronomy Methodology", False),
    ("6.4", "LLM Methodology and Guardrails", False),
    ("6.5", "Testing and Validation Methodology", False),
    ("7.", "Verification Performed", True),
    ("8.", "Known Limitations and Recommended Next Steps", True),
    ("9.", "Deployment and Migration Notes", True),
    ("10.", "Conclusion", True),
    ("A.", "Appendix A - API Endpoint Summary", True),
    ("B.", "Appendix B - Model Artifact Metadata", True),
    ("C.", "Appendix C - Environment Variables", True),
    ("D.", "Appendix D - File Map", True),
    ("E.", "Appendix E - Glossary", True),
]
for num, title, main in toc:
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(1)
    r = p.add_run(f"{num} {title}")
    r.bold = main
    if not main:
        p.paragraph_format.left_indent = Pt(20)
doc.add_page_break()

# ============================================================
# 1. EXECUTIVE SUMMARY
# ============================================================
doc.add_heading("1. Executive Summary", level=1)
para(
    "CropWise has evolved from a crop recommendation prototype into an AI-assisted crop advisory "
    "system for small and marginal farmers. The application still ranks suitable crops from soil and "
    "climate inputs, but it now also provides crop photos, a farmer-friendly Crop Guide, a global "
    "Gemini chat assistant, and computed growing, fertilizer and pest plans for each recommended crop."
)
para(
    "The most important design decision in Version 2.0 is the separation of responsibilities. Machine "
    "learning ranks crop suitability. A deterministic agronomy engine calculates dates, water advice, "
    "N/P/K deficits and fertilizer product quantities. Gemini is used only to explain already-computed "
    "plans or answer chat questions; it is never allowed to calculate or change a dose."
)
doc.add_paragraph("Key outcomes at a glance:")
table(
    ["Area", "Current result"],
    [
        ("Crop support", "51 crop classes in the prediction and plan APIs; 51 Crop Guide photo cards"),
        ("Training data", "6,596 records, 51 classes"),
        ("Model inputs", "13 engineered features + 5 FCM membership features = 18 model inputs"),
        ("Model performance", "99.92% holdout accuracy; 99.72% 5-fold CV (dataset-level, not field-level)"),
        ("Model artifact", "Random Forest champion, version 2, approximately 44 MB, trained 17 October 2026"),
        ("API surface", "20 routes including health, metadata and root; 17 domain routes"),
        ("Frontend", "React 18 + Vite + TypeScript + Tailwind, with Crop Guide, ChatWidget, CropPlanView and CropThumb"),
        ("Planning", "Deterministic growing, fertilizer and pest plans with field-specific values where data exists"),
        ("AI", "Server-side Gemini key, primary + fallback model, optional plan explanation, in-memory caches"),
        ("Verification", "51/51 crops returned 200 from /api/plan; chat and AI explanation tested; Playwright smoke tests passed; production build passes"),
    ],
)
doc.add_paragraph("")
para(
    "The current dataset is highly separable, so the 99.92% figure must not be presented as field-level "
    "accuracy. Independent field validation and agronomist review remain necessary before the plans are "
    "treated as prescriptive advice."
)

# ============================================================
# 2. CHANGES
# ============================================================
doc.add_heading("2. Changes Made in Version 2.0", level=1)

doc.add_heading("2.1 Frontend and User Experience", level=2)
table(
    ["Area", "Change", "Main files"],
    [
        ("Crop photos", "Added crop photos for all 51 supported crops, including bitter gourd and cabbage, plus a shared CropThumb component. Photos replace crop emoji only; UI icons remain unchanged. A safe fallback is used when an image is unavailable.", "frontend/public/crops/; src/components/CropThumb.tsx; src/utils/cropMeta.ts"),
        ("Crop Guide", "Added a farmer-facing /guide page with an 'In season now' strip, 'Understand your result' confidence-band explanations, a soil-numbers glossary, the existing visually unchanged clickable supported-crops grid (now all 51 crops), crop details on tap, a field selector, a plan modal and an AI explanation button. No technical model metrics or retrain UI were added.", "src/pages/CropGuide.tsx; src/components/CropPlanView.tsx"),
        ("Analytics replacement", "The farmer-facing Analytics page was replaced by Crop Guide. /analytics now redirects to /guide; the backend analytics endpoint remains available for model inspection.", "src/App.tsx; src/components/Layout.tsx"),
        ("Result plan modal", "Result now opens a full-width plan dialog for each recommended crop, with a per-crop in-memory cache and optional AI explanation.", "src/pages/Result.tsx"),
        ("History flow", "History entries can open Result and request a field-aware plan even when the stored history entry has no growing_guide.", "src/pages/History.tsx; src/pages/Result.tsx"),
        ("Chat assistant", "Added a floating ChatWidget on every authenticated app page. It is positioned below modals and sends recent conversation plus crop context.", "src/components/ChatWidget.tsx; src/components/Layout.tsx"),
        ("Crop labels", "Normalised display names from crop keys through a shared cropLabel helper (for example, sweet_potato -> Sweet Potato).", "src/utils/cropMeta.ts"),
        ("Source note", "Removed the cultivation-dataset source note from CropGuide and Result. The legacy text still exists in the API response but is not rendered.", "src/pages/CropGuide.tsx; src/pages/Result.tsx"),
        ("Technical jargon", "Removed ML-pipeline jargon from the Result footer and sidebar; replaced it with farmer-facing wording such as Practical crop guidance.", "src/pages/Result.tsx; src/components/Layout.tsx"),
    ],
)

doc.add_heading("2.2 Backend, API and Services", level=2)
table(
    ["Area", "Change", "Main files"],
    [
        ("Personalized plan API", "Added POST /api/plan. It accepts a crop, optional field, optional N/P/K and climate values, soil type and include_ai. Unknown crops or fields return 404.", "app/routers/plan.py; app/schemas/plan.py"),
        ("Deterministic plan engine", "Added the growing/fertilizer/pest plan engine. It computes sowing and harvest windows, irrigation advice, soil preparation, risks, N/P/K targets or deficits, DAP/urea/MOP quantities, stage splits and pests.", "app/services/crop_plan.py"),
        ("Crop profiles", "Added structured profiles for 51 crops: duration, planting windows, relative water need, critical stages and perennial flag. October is treated as Rabi in the season logic.", "app/services/crop_profiles.py"),
        ("Agronomy knowledge", "Added crop N/P/K targets, dose units, fertilizer stage schedules and pest guides for the supported crops.", "app/services/agronomy.py"),
        ("Weather and soil", "Added cached Open-Meteo 7-day forecast and ISRIC SoilGrids topsoil pH/clay/sand clients. Both fail soft; SoilGrids never supplies N/P/K.", "app/services/field_data.py"),
        ("AI narration", "Added plan explanation that receives only the computed plan. Explanations are cached by plan-content hash; the model is instructed not to change any number.", "app/services/plan_narrative.py"),
        ("Shared LLM client", "Added an httpx-based Gemini client using the OpenAI-compatible endpoint, with primary and fallback models, timeout handling and a friendly unavailable state.", "app/services/llm.py"),
        ("Chat API", "Added POST /api/chat with message validation, crop context, agronomy facts, simple-language rules and server-side key handling.", "app/routers/chat.py; app/schemas/chat.py"),
        ("Recommendation integration", "/api/recommend now predicts top-5 candidates, re-ranks to top-3 by season and soil type, builds a growing guide for each crop and returns the original inputs for the Result plan.", "app/routers/recommend.py"),
        ("Legacy guidance", "/api/guidance remains available for compatibility, but the UI uses /api/plan for personalized plans.", "app/routers/guidance.py"),
        ("No retrain UI", "Model status and retrain remain API-only; no farmer-facing retrain controls were added.", "app/routers/model.py"),
    ],
)

doc.add_heading("2.3 Machine Learning Model", level=2)
table(
    ["Item", "Version 2.0 state"],
    [
        ("Dataset", "6,596 rows in backend/data/crop_data.csv"),
        ("Classes", "51 crop classes"),
        ("Engineered features", "N, P, K, temperature, humidity, pH, rainfall + NPK_sum, N_P_ratio, K_P_ratio, log_rainfall, temp_humidity_index, aridity_index = 13 features"),
        ("Cluster features", "5 fuzzy membership features from 5 FCM clusters, appended after scaling and PSO weighting"),
        ("Pipeline", "StandardScaler -> PSO feature weights -> FCM membership -> classifier"),
        ("Classifier", "Random Forest champion selected over HistGradientBoosting by cross-validation"),
        ("Evaluation", "80/20 holdout with the split made before fitting; 5-fold cross-validation"),
        ("Holdout accuracy", "99.92% (0.99924)"),
        ("Cross-validation", "99.72% mean (0.99716), standard deviation approximately 0.17%"),
        ("Artifact", "backend/app/ml/artifacts/model.joblib, approximately 44 MB; model_v2.joblib and manifest.json also present"),
        ("Trained", "2026-10-17T12:59:49Z, manifest version 2"),
        ("Caveat", "These are dataset-level metrics. The current dataset is highly separable; field-level accuracy has not been established."),
    ],
)

doc.add_heading("2.4 Database and Data Handling", level=2)
bullet("Field soil values: ", "fields now store soil_n, soil_p, soil_k, soil_ph and soil_tested_at. Startup performs an additive SQLite migration for these columns.")
bullet("History: ", "recommendations store the original N/P/K and climate inputs, top crops, actual crop/yield and feedback fields.")
bullet("Feedback: ", "feedback is stored on the recommendation row; there is no separate feedback table.")
bullet("Plans and chat: ", "generated plans and chat messages are not persisted. Plans are computed on demand; chat lives in the browser session.")
bullet("No new tables: ", "Version 2.0 added no database tables. Existing fields and recommendations tables were extended only where needed.")

doc.add_heading("2.5 Navigation, Content and Compatibility", level=2)
bullet("Navigation: ", "sidebar now exposes Dashboard, New Recommendation, Fields, History, Crop Guide and Profile. Analytics is no longer a farmer-facing page.")
bullet("Routes: ", "/login, /, /recommend, /result, /fields, /history, /guide, /profile. /analytics redirects to /guide.")
bullet("Auth: ", "login is a demo Google-style sign-in stored in localStorage. It is not production OAuth and has no backend session or JWT.")
bullet("Compatibility: ", "the legacy /api/analytics, /api/guidance and model endpoints remain for compatibility and inspection.")
bullet("Crop coverage: ", "the backend supports 51 crops and the Crop Guide now displays all 51 photo cards. Bitter gourd and cabbage were added to CROP_META with their own photos, so the grid count and the model/API count match.")
bullet("Source note: ", "the legacy source note remains in the API but is hidden from all UI cards.")

# ============================================================
# 3. SOFTWARE REQUIREMENTS
# ============================================================
doc.add_heading("3. Updated Software Requirements", level=1)
para(
    "The runtime remains a Python FastAPI backend and a React/Vite frontend. Version 2.0 adds outbound "
    "HTTP calls to Open-Meteo, ISRIC SoilGrids and Google Gemini. No new frontend package was required "
    "for the chat or plan UI."
)

doc.add_heading("3.1 Backend Runtime", level=2)
table(
    ["Component", "Requirement", "Notes"],
    [
        ("Python", "3.11+ recommended", "Tested on 3.13.3; Docker uses python:3.11-slim"),
        ("FastAPI", ">= 0.109", "Tested 0.141.1"),
        ("uvicorn[standard]", ">= 0.29", "ASGI server"),
        ("SQLAlchemy", ">= 2.0", "Tested 2.0.52"),
        ("Pydantic", ">= 2.6", "Tested 2.13.5"),
        ("pydantic-settings", ">= 2.1", "Environment configuration"),
        ("pandas", ">= 2.0", "Tested 2.2.3"),
        ("scikit-learn", ">= 1.4", "Tested 1.6.1"),
        ("NumPy", ">= 1.26", "Tested 2.2.5"),
        ("Joblib", ">= 1.3", "Model serialization"),
        ("httpx", ">= 0.27, < 0.28", "Tested 0.27.2; weather, soil and Gemini calls"),
        ("psycopg2-binary", ">= 2.9", "PostgreSQL driver"),
        ("python-multipart", ">= 0.0.9", "Form/multipart support"),
        ("SciPy", "Transitive via scikit-learn", "No direct application import"),
        ("python-docx", ">= 1.1", "Report generation only; not a runtime dependency"),
    ],
)

doc.add_heading("3.2 Frontend Runtime", level=2)
table(
    ["Component", "Requirement", "Notes"],
    [
        ("Node.js", "20 LTS+ recommended", "Tested 24.11.1; Docker uses node:20-alpine"),
        ("npm", ">= 10", "Tested 11.6.2"),
        ("React", "18.3.1", "React DOM 18.3.1"),
        ("React Router DOM", "6.23.0", "Client-side routing"),
        ("Vite", "5.2.12+", "Tested 5.4.21"),
        ("TypeScript", "5.4.5", "Type checking in build"),
        ("Tailwind CSS", "3.4.3", "UI styling"),
        ("Zustand", "4.5.2", "Global app state"),
        ("Chart.js", "4.4.2", "Charts where used"),
        ("react-chartjs-2", "5.2.0", "React chart wrapper"),
        ("Leaflet", "1.9.4", "Maps"),
        ("react-leaflet", "4.2.1", "React map wrapper"),
        ("@headlessui/react", "1.7.19", "Accessible dialogs and menus"),
        ("Playwright", "Optional", "End-to-end smoke testing; not a runtime dependency"),
    ],
)

doc.add_heading("3.3 Deployment and External Services", level=2)
table(
    ["Component", "Requirement", "Notes"],
    [
        ("SQLite", "Development default", "crop_recommend.db; automatic table creation"),
        ("PostgreSQL", "16", "Production/Docker via DATABASE_URL"),
        ("Docker + Docker Compose", "Optional", "postgres, backend and frontend services"),
        ("Nginx", "Production frontend", "Serves the Vite build and proxies /api"),
        ("Git", "Version control", "Repository and change history"),
        ("Open-Meteo", "Free, no API key", "7-day forecast: rain, ET0, max/min temperature"),
        ("ISRIC SoilGrids", "Free, no API key", "Modelled 0-5 cm pH, clay and sand; never N/P/K"),
        ("Google Gemini", "API key required", "OpenAI-compatible endpoint; key stays server-side"),
        ("Browser", "Current Chrome, Edge, Firefox or Safari", "JavaScript enabled; internet for live data and AI"),
        ("Operating system", "Windows 10/11, macOS 12+, or Linux", "Development verified on Windows"),
    ],
)

doc.add_heading("3.4 Gemini Configuration", level=2)
table(
    ["Variable", "Default", "Purpose"],
    [
        ("LLM_API_KEY", "Empty", "Server-side Gemini key. Empty disables chat and AI narration gracefully."),
        ("LLM_MODEL", "gemini-3.8-flash", "Primary model"),
        ("LLM_FALLBACK_MODEL", "gemini-flash-lite-latest", "Fallback when the primary model is unavailable"),
        ("LLM_BASE_URL", "https://generativelanguage.googleapis.com/v1beta/openai", "OpenAI-compatible Gemini endpoint"),
    ],
)
para(
    "Security note: the key must be supplied through the backend environment or .env file and must never "
    "be included in frontend code, browser requests, screenshots or committed files."
)

# ============================================================
# 4. HARDWARE REQUIREMENTS
# ============================================================
doc.add_heading("4. Updated Hardware Requirements", level=1)
table(
    ["Specification", "Minimum", "Recommended"],
    [
        ("Processor", "2 GHz dual-core", "4-core or better"),
        ("RAM", "4 GB", "8 GB or more"),
        ("Storage", "2 GB free", "5 GB+ SSD free"),
        ("Display", "1366 x 768", "1920 x 1080"),
        ("Network", "Broadband internet", "Broadband, low latency"),
        ("GPU", "Not required", "Not required"),
        ("Production server", "2 vCPU / 4 GB RAM / 10 GB disk", "4 vCPU / 8 GB RAM / 20 GB SSD"),
    ],
)
doc.add_paragraph("")
para("Why these values are sufficient:")
bullet("Model: ", "the serialized model is approximately 44 MB. Training uses 6,596 rows and 13 engineered features; it is CPU-friendly and does not need a GPU.")
bullet("Images: ", "the 51 crop photos total approximately 7.2 MB.")
bullet("Dataset: ", "crop_data.csv is approximately 373 KB.")
bullet("Dependencies: ", "the Python environment, node_modules, Vite build, Docker images and PostgreSQL volume together make 5 GB of free space a sensible recommendation.")
bullet("Frontend: ", "the production JavaScript bundle is approximately 603 KB and CSS approximately 44 KB; the app is lightweight in the browser.")
bullet("Offline behavior: ", "plans still work without Open-Meteo or SoilGrids because external calls fail soft. Internet is required for live data and Gemini features.")
bullet("Production: ", "PostgreSQL, Docker and Nginx add overhead; use the higher production recommendation for a shared deployment.")

# ============================================================
# 5. ARCHITECTURE
# ============================================================
doc.add_heading("5. Current System Architecture", level=1)

doc.add_heading("5.1 Architectural Overview", level=2)
para(
    "CropWise is a layered web application with a separate intelligence layer. The design is hybrid: "
    "ML handles crop suitability, a deterministic engine handles agronomy, and an LLM handles language."
)
code(
    "+---------------------------------------------------------------+\n"
    "| Presentation: React 18 + Vite + TypeScript + Tailwind          |\n"
    "| Dashboard | Recommend | Result | Crop Guide | Fields | History |\n"
    "| Profile | CropThumb | CropPlanView | ChatWidget                |\n"
    "+-----------------------------+---------------------------------+\n"
    "                              | REST / JSON\n"
    "+-----------------------------v---------------------------------+\n"
    "| API: FastAPI + Pydantic + SQLAlchemy                           |\n"
    "| /api/recommend | /api/plan | /api/chat | /api/fields ...      |\n"
    "+--------------+-----------------------------+------------------+\n"
    "               |                             |\n"
    "+--------------v-----------------+  +--------v-----------------+\n"
    "| Intelligence                   |  | Data                     |\n"
    "| ML ranking | Plan engine | LLM  |  | SQLite / PostgreSQL      |\n"
    "| FCM+PSO+RF | deterministic | Gemini | fields, recommendations |\n"
    "+--------------+-----------------+  +--------------------------+\n"
    "               |\n"
    "+--------------v-----------------------------------------------+\n"
    "| External: Open-Meteo | ISRIC SoilGrids | Google Gemini         |\n"
    "+--------------------------------------------------------------+"
)

doc.add_heading("5.2 Layer Responsibilities", level=2)
para("Presentation layer")
bullet("Pages: ", "Login, Dashboard, New Recommendation, Result, Fields, History, Crop Guide and Profile.")
bullet("Shared components: ", "Layout, CropThumb, CropPlanView, ChatWidget, Cards, Badges, Buttons, Inputs and Accordion.")
bullet("State and presentation: ", "Zustand for app state, Tailwind for styling, Chart.js for charts and Leaflet for maps.")
bullet("Crop Guide: ", "the main farmer-facing guidance page; in-season crops, confidence bands, soil glossary, supported-crop grid and plan modal.")
bullet("Result: ", "crop cards plus a full-width plan modal, per-crop cache and optional AI explanation.")

para("API layer")
bullet("Routers: ", "recommend, plan, chat, weather, fields, history, feedback, guidance, analytics and model.")
bullet("Validation: ", "Pydantic schemas enforce ranges for N/P/K, pH, temperature, humidity, rainfall, field IDs and chat message length.")
bullet("Security headers: ", "X-Content-Type-Options: nosniff and X-Frame-Options: DENY.")
bullet("CORS: ", "configured from CORS_ORIGINS; development defaults include localhost:5173 and localhost:3000.")
bullet("Authentication: ", "currently a demo frontend sign-in. There is no backend JWT or OAuth session.")

para("Intelligence layer")
bullet("ML ranking: ", "FCM + PSO + Random Forest produces top-5 candidates; contextual re-ranking returns top-3. The model does not calculate doses.")
bullet("Plan engine: ", "resolves soil values (explicit request > field soil test > SoilGrids pH), fetches weather, computes dates/water/N/P/K and returns growing, fertilizer and pest sections.")
bullet("LLM: ", "shared Gemini client; chat answers are grounded with crop facts and agronomy snippets; plan narration only explains computed values.")
bullet("Caching: ", "weather 1 hour, SoilGrids 7 days, AI explanations by plan-content hash with a 300-entry cap, and per-crop Result cache in the browser.")

para("Data layer")
bullet("Tables: ", "fields and recommendations. Feedback fields are stored on recommendations.")
bullet("Field columns: ", "id, name, latitude, longitude, area_hectares, soil_type, soil_n, soil_p, soil_k, soil_ph, soil_tested_at, created_at.")
bullet("Recommendation columns: ", "id, field_id, N, P, K, temperature, humidity, ph, rainfall, top_crops JSON, actual_crop, actual_yield, feedback_rating, feedback_notes, created_at.")
bullet("No new tables: ", "plans and chat messages are not persisted; plans are computed on demand.")
bullet("Schema creation: ", "SQLAlchemy metadata.create_all plus an additive SQLite migration for new field soil columns.")

para("External services")
bullet("Open-Meteo: ", "7-day forecast for the field's latitude/longitude; returns rain, ET0 and temperature; no key.")
bullet("ISRIC SoilGrids: ", "modelled 0-5 cm pH, clay and sand; no key; not a source of N/P/K.")
bullet("Google Gemini: ", "OpenAI-compatible endpoint; primary and fallback models; key server-side only.")

doc.add_heading("5.3 Data Flows", level=2)
para("Recommendation flow:")
code(
    "1. Farmer submits N, P, K, temperature, humidity, pH, rainfall and optional field/soil type.\n"
    "2. Frontend sends POST /api/recommend.\n"
    "3. Backend validates the request and derives the 13 engineered features.\n"
    "4. Model predicts top-5 candidates; rerank applies season and soil-type adjustments.\n"
    "5. Top-3 crops are saved to recommendation history.\n"
    "6. Plan engine builds a growing guide for each top crop.\n"
    "7. Response returns top crops, growing guide, limiting factor and original inputs."
)
para("Personalized plan flow:")
code(
    "1. Crop Guide or Result requests POST /api/plan with a crop and optional field/inputs.\n"
    "2. Backend validates the crop and field.\n"
    "3. If the field has coordinates, weather and (when needed) SoilGrids run in parallel.\n"
    "4. Soil values resolve in priority order: explicit values > field soil test > SoilGrids pH.\n"
    "5. The deterministic engine builds growing, fertilizer and pest sections.\n"
    "6. If include_ai=true, Gemini narrates the already-computed plan and caches the result.\n"
    "7. Response returns the plan; external-service failures leave the plan usable."
)
para("Chat flow:")
code(
    "1. ChatWidget sends recent messages plus optional crop context to POST /api/chat.\n"
    "2. Backend validates message count, length and roles.\n"
    "3. System prompt and agronomy crop facts are assembled server-side.\n"
    "4. Gemini primary model is called; fallback model is tried if needed.\n"
    "5. A friendly reply is returned; the key never reaches the browser."
)

doc.add_heading("5.4 Caching, Failure Handling and Security", level=2)
bullet("Weather cache: ", "1 hour per rounded latitude/longitude.")
bullet("Soil cache: ", "7 days per rounded latitude/longitude.")
bullet("Plan explanation cache: ", "SHA-256 of plan content; maximum 300 entries; repeated card views avoid new quota use.")
bullet("Result cache: ", "per-crop plan results are held in component memory for the current session.")
bullet("Fail-soft: ", "Open-Meteo or SoilGrids failure returns no external data, and the plan still uses saved values or standard targets.")
bullet("AI failure: ", "LLM unavailability returns no AI explanation; the deterministic plan remains available.")
bullet("Secrets: ", "Gemini key is server-side; CORS is restricted; security headers are applied; HTTPS is recommended in production.")
bullet("Privacy: ", "field coordinates are sent to weather/soil services; plan text is sent to Gemini for narration; user identity and chat history are not intentionally sent or persisted.")
bullet("Auth caveat: ", "demo sign-in is not production authentication and must be replaced before a public deployment.")

# ============================================================
# 6. METHODOLOGY
# ============================================================
doc.add_heading("6. Current Methodology", level=1)
para(
    "Version 2.0 uses a hybrid methodology: iterative software delivery, CRISP-DM for machine learning, "
    "deterministic agronomy for decisions, grounded LLM generation for language, and layered verification."
)

doc.add_heading("6.1 Development Methodology", level=2)
bullet("Approach: ", "iterative/Agile delivery in small vertical slices, with frontend and backend changes verified together.")
bullet("Design principles: ", "separate explanation from calculation; keep computed plans reproducible; fail soft on external services; keep the farmer UI free of ML jargon.")
bullet("Change control: ", "the working tree is uncommitted; the model artifact is versioned through manifest.json.")
bullet("Build discipline: ", "TypeScript + Vite production build, backend startup/health checks and Playwright smoke tests.")

doc.add_heading("6.2 Machine Learning Methodology", level=2)
para("The ML component follows CRISP-DM:")
doc.add_paragraph("1. Business understanding: identify crops suited to a soil and climate profile.", style="List Number")
doc.add_paragraph("2. Data understanding: 6,596 records across 51 crop classes.", style="List Number")
doc.add_paragraph("3. Data preparation: derive 13 engineered features, split 80/20 before fitting, and avoid leakage.", style="List Number")
doc.add_paragraph("4. Modeling: StandardScaler -> PSO feature weights -> 5-cluster FCM membership -> Random Forest vs HistGradientBoosting champion selected by CV.", style="List Number")
doc.add_paragraph("5. Evaluation: holdout accuracy, 5-fold CV, confusion matrix, classification report and feature importance.", style="List Number")
doc.add_paragraph("6. Deployment: serialize the champion, load it lazily in the recommendation router, and fall back to a heuristic if it cannot load.", style="List Number")
para(
    "The model returns top-5 candidates; rerank.py applies season and soil-type adjustments and returns top-3. "
    "October is treated as Rabi in the season logic. The 99.92% holdout and 99.72% CV are dataset-level results "
    "and require field validation before being treated as real-world accuracy."
)

doc.add_heading("6.3 Agronomy Methodology", level=2)
bullet("Knowledge source: ", "structured crop profiles and reference N/P/K targets derived from the curated cultivation dataset and standard package-of-practices values.")
bullet("Growing plan: ", "month-based planting window, duration range, harvest estimate, soil preparation, irrigation advice, risk notes and critical stages.")
bullet("Field data: ", "7-day forecast and modelled pH/texture when a field has coordinates; saved soil-test values otherwise.")
bullet("Fertilizer: ", "standard crop target when no soil N/P/K exists; target minus available values when a soil test exists. Missing values are clearly flagged.")
bullet("Products: ", "DAP (18-46-0), urea (46-0-0) and MOP (0-0-60), with kg/ha, kg/acre and 50 kg bags/acre. P and K targets are treated as P2O5 and K2O.")
bullet("Stages: ", "reference stage shares for nitrogen; P and K are currently assumed basal. DAP nitrogen is credited against total urea.")
bullet("Per-plant crops: ", "banana, fruit trees and other per-plant/per-tree crops skip the per-hectare product conversion and are labelled accordingly.")
bullet("Pests: ", "crop-specific static pest/treatment/organic guidance; not a live integrated pest-management model.")
bullet("Safety: ", "product grades and doses must be confirmed with a local dealer or agriculture officer; plans are decision support, not a prescription.")
bullet("No AI doses: ", "Gemini never calculates or changes fertilizer quantities.")

doc.add_heading("6.4 LLM Methodology and Guardrails", level=2)
bullet("Grounding: ", "chat receives crop facts and agronomy snippets; plan narration receives the computed plan only.")
bullet("Guardrails: ", "simple English, short paragraphs, no ML jargon, no invented numbers, no changed doses, and a recommendation to confirm products locally.")
bullet("Model policy: ", "gemini-3.8-flash primary, gemini-flash-lite-latest fallback, 2048 max tokens, timeout handling and a safe unavailable response.")
bullet("Caching: ", "plan explanations are cached by content hash; chat is not cached server-side.")
bullet("Optionality: ", "AI explanation is requested with include_ai and is not part of the deterministic dose calculation.")
bullet("Latency: ", "AI calls add roughly 7-13 seconds; the base plan is available without AI.")

doc.add_heading("6.5 Testing and Validation Methodology", level=2)
bullet("API regression: ", "all 51 crops exercised against /api/plan; recommendation and chat endpoints exercised.")
bullet("Scenario testing: ", "no field, field with coordinates, field with soil test, per-plant banana, explicit N/P/K and include_ai.")
bullet("Frontend testing: ", "Crop Guide standard/field plans, Result modal, history-to-result flow and AI explanation via Playwright smoke tests.")
bullet("Build verification: ", "TypeScript + Vite production build.")
bullet("Manual review: ", "plan wording, source-note removal, crop labels and farmer-facing language.")
bullet("Future validation: ", "agronomist review, independent field trials, soil-test calibration and regional package-of-practices validation.")

# ============================================================
# 7. VERIFICATION
# ============================================================
doc.add_heading("7. Verification Performed", level=1)
table(
    ["Check", "Result", "Evidence / note"],
    [
        ("Backend health", "PASS", "/health returns 200"),
        ("All crop plans", "PASS", "51/51 crops returned 200 from /api/plan"),
        ("No-field plan", "PASS", "Standard crop dose shown and clearly labelled"),
        ("Field with coordinates", "PASS", "Live Open-Meteo forecast used; plan returns personalized basis"),
        ("Field with soil test", "PASS", "Field-specific N/P/K deficit calculated"),
        ("Per-plant banana", "PASS", "Per-hectare product conversion skipped"),
        ("Explicit N/P/K", "PASS", "Request values take priority over stored values"),
        ("include_ai=true", "PASS", "AI explanation returned; plan content unchanged"),
        ("Recommendation integration", "PASS", "Top-3 crops returned with computed growing guides and inputs"),
        ("Chat API", "PASS", "Returned a reply in approximately 7.9 seconds"),
        ("AI explanation", "PASS", "Complete 949-character explanation in the verification run"),
        ("Crop Guide UI", "PASS", "In-season strip, glossary, 51-crop grid, photos and standard/field plan modals"),
        ("Result plan modal", "PASS", "Full-width dialog, cache and AI explain button"),
        ("History flow", "PASS", "History -> Result -> field plan works without stored growing_guide"),
        ("Analytics redirect", "PASS", "/analytics redirects to /guide"),
        ("Source note removal", "PASS", "No source note rendered in Crop Guide or Result"),
        ("Frontend build", "PASS", "Production build passes"),
        ("Playwright smoke suite", "PASS", "Crop Guide, Result, history and AI explanation flows"),
    ],
)
doc.add_paragraph("")
para("Known warnings and caveats:")
bullet("Vite bundle warning: ", "the production build reports a chunk larger than 500 kB. This warning is pre-existing and non-blocking.")
bullet("Dataset accuracy: ", "99.92% holdout and 99.72% CV are not field accuracy. The current dataset is highly separable.")
bullet("External outage: ", "fail-soft behavior is designed and verified by code inspection; a controlled outage test is still recommended.")
bullet("Crop coverage: ", "51 Crop Guide photo cards for 51 supported crops; bitter gourd and cabbage are now included.")
bullet("Auth: ", "demo sign-in only; not production authentication.")

# ============================================================
# 8. LIMITATIONS / NEXT STEPS
# ============================================================
doc.add_heading("8. Known Limitations and Recommended Next Steps", level=1)
table(
    ["Limitation", "Impact", "Recommended next step"],
    [
        ("Dataset-level accuracy only", "High reported accuracy may not transfer to real farms", "Run independent field validation and agronomist review"),
        ("DAP/MOP assumed basal", "Stage-level P/K advice is simplified", "Implement exact stage-level N/P/K allocation"),
        ("DAP N credit not reallocated", "Nitrogen stage shares may be slightly uneven", "Allocate DAP nitrogen across N stages"),
        ("Missing N/P/K treated as 0 in some deficit paths", "Incomplete soil tests can overstate deficits", "Show a clear incomplete-test state and avoid zero-filling"),
        ("Per-plant targets lack kg/plant conversion", "Banana/tree crops receive reference guidance rather than a full product plan", "Add per-plant and per-tree dose ranges from the cultivation CSV"),
        ("Pest guide is static and generic", "No economic threshold, PHI or regulatory checks", "Build an IPM knowledge base with ETL, PHI and local validation"),
        ("Month-based crop calendar", "Variety and region differences are not captured", "Add regional calendars and weather-driven water balance"),
        ("No regional soil-test calibration", "Targets are reference values", "Add regional calibration and organic-matter credit"),
        ("Demo authentication", "Not suitable for public deployment", "Add real OAuth/JWT sessions and authorization"),
        ("No chat/plan persistence", "Users cannot revisit conversations or saved plans", "Add optional plan/chat history with privacy controls"),
    ],
)

# ============================================================
# 9. DEPLOYMENT
# ============================================================
doc.add_heading("9. Deployment and Migration Notes", level=1)
para("Minimum deployment sequence:")
doc.add_paragraph("1. Create backend/.env from backend/.env.example and set LLM_API_KEY.", style="List Number")
doc.add_paragraph("2. Install backend dependencies: pip install -r requirements.txt.", style="List Number")
doc.add_paragraph("3. Train or confirm the model artifact: python train.py (optional if model.joblib is present).", style="List Number")
doc.add_paragraph("4. Start the backend from backend/: python -m uvicorn app.main:app --host 0.0.0.0 --port 8000.", style="List Number")
doc.add_paragraph("5. Install frontend dependencies: cd frontend && npm install.", style="List Number")
doc.add_paragraph("6. Build and serve: npm run build; for development, npm run dev.", style="List Number")
doc.add_paragraph("7. Verify /health, /api/meta, /api/recommend, /api/plan and /api/chat.", style="List Number")
doc.add_paragraph("8. For Docker: docker compose up --build from the project root.", style="List Number")
doc.add_paragraph("")
para("Database migration: ", "startup calls Base.metadata.create_all and an additive SQLite migration that adds soil_n, soil_p, soil_k, soil_ph and soil_tested_at to fields. No manual migration is required for the current SQLite development database. Production PostgreSQL should be reviewed before deployment.")
para("Backend restart: ", "the backend is currently run without --reload. Restart it after backend changes, especially after changing services, routes or environment variables.")
para("Frontend routing: ", "use http://localhost:5173 in development. Vite is bound in a way that can make 127.0.0.1:5173 fail on some machines; localhost is the reliable development URL.")

# ============================================================
# 10. CONCLUSION
# ============================================================
doc.add_heading("10. Conclusion", level=1)
para(
    "Version 2.0 turns CropWise into a more complete farmer-facing advisory system. The original ML "
    "recommendation flow remains, but it is now surrounded by a Crop Guide, real crop photos, a chat "
    "assistant and field-aware growing, fertilizer and pest plans. The most important architectural "
    "change is the separation of deterministic agronomy from LLM narration: the numbers are computed by "
    "the application, and the model only explains them."
)
para(
    "The implementation is functional and verified at the API, UI and build levels. The remaining work is "
    "mainly accuracy and production readiness: field validation, stage-level nutrient allocation, per-plant "
    "doses, a validated IPM knowledge base, regional calibration and real authentication. Those items are "
    "recorded above rather than hidden behind the current dataset-level accuracy."
)

# ============================================================
# APPENDICES
# ============================================================
doc.add_heading("Appendix A - API Endpoint Summary", level=1)
table(
    ["Endpoint", "Method", "Purpose"],
    [
        ("/", "GET", "API root message"),
        ("/health", "GET", "Health check"),
        ("/api/meta", "GET", "System name, version and pipeline metadata"),
        ("/api/recommend", "POST", "Soil/climate input -> top-3 crops, growing guide and inputs"),
        ("/api/plan", "POST", "Crop and optional field -> growing, fertilizer, pest and optional AI plan"),
        ("/api/chat", "POST", "Grounded chatbot reply"),
        ("/api/weather", "POST", "Latitude/longitude -> climate averages"),
        ("/api/guidance", "POST", "Legacy guidance endpoint; UI uses /api/plan"),
        ("/api/analytics", "GET", "Model accuracy, feature weights and confusion matrix"),
        ("/api/fields", "GET", "List fields"),
        ("/api/fields", "POST", "Create field"),
        ("/api/fields/{id}", "GET", "Get one field"),
        ("/api/fields/{id}", "PUT", "Update field"),
        ("/api/fields/{id}", "DELETE", "Delete field"),
        ("/api/history", "GET", "Recommendation history"),
        ("/api/history/{id}", "PATCH", "Link a history entry to a field"),
        ("/api/feedback", "POST", "Save actual crop, yield, rating and notes"),
        ("/api/feedback/{id}", "GET", "Read saved feedback"),
        ("/api/model/status", "GET", "Model artifact status"),
        ("/api/model/retrain", "POST", "Retrain the model (API only; no farmer UI)"),
    ],
)

doc.add_heading("Appendix B - Model Artifact Metadata", level=1)
code(
    "Location        : backend/app/ml/artifacts/model.joblib\n"
    "Also present    : model_v2.joblib, manifest.json\n"
    "Format          : Joblib serialized Python object\n"
    "Size            : approximately 44 MB\n"
    "Artifact version: 2\n"
    "Trained at      : 2026-10-17T12:59:49Z\n"
    "Dataset         : 6,596 rows, 51 classes\n"
    "Engineered      : 13 features (7 raw + 6 derived)\n"
    "Cluster inputs  : 5 FCM membership features\n"
    "Model           : Random Forest champion\n"
    "Holdout accuracy: 99.92%\n"
    "5-fold CV       : 99.72% mean, approximately 0.17% standard deviation\n"
    "Data source     : backend/data/crop_data.csv"
)

doc.add_heading("Appendix C - Environment Variables", level=1)
table(
    ["Variable", "Required", "Default / example", "Purpose"],
    [
        ("DATABASE_URL", "No", "sqlite:///./crop_recommend.db", "SQLAlchemy database URL"),
        ("SECRET_KEY", "Production", "dev-secret-change-in-prod", "Application secret; replace in production"),
        ("CORS_ORIGINS", "No", "http://localhost:5173,http://localhost:3000", "Allowed frontend origins"),
        ("LLM_API_KEY", "For AI", "empty", "Gemini key; server-side only"),
        ("LLM_MODEL", "No", "gemini-3.8-flash", "Primary Gemini model"),
        ("LLM_FALLBACK_MODEL", "No", "gemini-flash-lite-latest", "Fallback Gemini model"),
        ("LLM_BASE_URL", "No", "https://generativelanguage.googleapis.com/v1beta/openai", "OpenAI-compatible endpoint"),
        ("VITE_API_URL", "No", "empty (use Vite proxy)", "Optional frontend API base URL"),
    ],
)

doc.add_heading("Appendix D - File Map", level=1)
code(
    "backend/app/routers/plan.py          POST /api/plan\n"
    "backend/app/routers/chat.py          POST /api/chat\n"
    "backend/app/routers/recommend.py     recommendation + computed growing guide\n"
    "backend/app/services/crop_plan.py    deterministic plan engine\n"
    "backend/app/services/crop_profiles.py 51 crop profiles\n"
    "backend/app/services/agronomy.py     NPK targets, stages, pests\n"
    "backend/app/services/field_data.py   Open-Meteo + SoilGrids caches\n"
    "backend/app/services/plan_narrative.py  Gemini plan explanation + cache\n"
    "backend/app/services/llm.py          shared Gemini client\n"
    "backend/app/schemas/plan.py          plan request/response models\n"
    "backend/app/schemas/chat.py          chat request/response models\n"
    "frontend/src/pages/CropGuide.tsx     Crop Guide page\n"
    "frontend/src/pages/Result.tsx        Result + plan modal\n"
    "frontend/src/components/CropPlanView.tsx  shared plan renderer\n"
    "frontend/src/components/ChatWidget.tsx    global chat widget\n"
    "frontend/src/components/CropThumb.tsx     crop photo with fallback\n"
    "frontend/src/utils/cropMeta.ts       crop names, photos, labels\n"
    "frontend/public/crops/               51 crop images\n"
)

doc.add_heading("Appendix E - Glossary", level=1)
table(
    ["Term", "Meaning"],
    [
        ("N, P, K", "Nitrogen, phosphorus and potassium; major plant nutrients"),
        ("NPK", "Combined nitrogen, phosphorus and potassium"),
        ("DAP", "Di-ammonium phosphate, standard grade 18-46-0"),
        ("MOP", "Muriate of potash, standard grade 0-0-60"),
        ("Urea", "Nitrogen fertilizer, standard grade 46-0-0"),
        ("kg/ha, kg/acre", "Product quantity per hectare or per acre"),
        ("P2O5, K2O", "Standard oxide forms used when expressing phosphorus and potassium targets"),
        ("ET0", "Reference evapotranspiration; used to frame irrigation advice"),
        ("SoilGrids", "ISRIC modelled soil property service; pH/clay/sand only"),
        ("Kharif/Rabi/Zaid", "Major Indian cropping seasons"),
        ("FCM", "Fuzzy C-Means clustering"),
        ("PSO", "Particle Swarm Optimization"),
        ("RF", "Random Forest"),
        ("CRISP-DM", "Cross-Industry Standard Process for Data Mining"),
        ("LLM", "Large language model; here, Gemini used for chat and narration"),
        ("Fail-soft", "Continue with saved/standard values when an external service is unavailable"),
    ],
)

# ============================================================
# SAVE
# ============================================================
doc.save(OUT)
print(f"Document saved to: {OUT}")
